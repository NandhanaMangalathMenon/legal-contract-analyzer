from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument

from agent1.ocr.interface import OCREngine, StubOCREngine
from agent1.ocr.gemini_engine import GeminiOCREngine
from core.exceptions import IngestionError
from core.ids import document_hash


@dataclass
class PageText:
    page: int
    text: str


@dataclass
class IngestedDocument:
    filename: str
    file_type: str
    pages: list[PageText]
    raw_bytes: bytes
    content_hash: str
    ocr_used: bool = False
    extraction_confidence: float = 1.0
    warnings: list[str] | None = None

    @property
    def full_text(self) -> str:
        return "\n\n".join(page.text for page in self.pages)


def ingest_document(path: Path, ocr: OCREngine | None = None) -> IngestedDocument:
    path = Path(path)
    if not path.exists():
        raise IngestionError(f"File not found: {path}")
    suffix = path.suffix.lower()
    raw = path.read_bytes()
    ocr = ocr or GeminiOCREngine()
    if suffix == ".pdf":
        return _ingest_pdf(path, raw, ocr)
    if suffix == ".docx":
        return _ingest_docx(path, raw)
    if suffix in {".txt", ".md"}:
        return _ingest_txt(path, raw)
    raise IngestionError(f"Unsupported file type: {suffix}")


def _ingest_txt(path: Path, raw: bytes) -> IngestedDocument:
    text = raw.decode("utf-8", errors="replace")
    pages = _split_page_markers(text)
    return IngestedDocument(
        filename=path.name,
        file_type="txt",
        pages=pages,
        raw_bytes=raw,
        content_hash=document_hash(raw),
    )


def _ingest_docx(path: Path, raw: bytes) -> IngestedDocument:
    document = DocxDocument(str(path))
    paragraphs = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    tables: list[str] = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            tables.append("TABLE:\n" + "\n".join(rows))
    text = "\n".join(paragraphs + tables)
    pages = _split_page_markers(text) if text else [PageText(page=1, text="")]
    return IngestedDocument(
        filename=path.name,
        file_type="docx",
        pages=pages,
        raw_bytes=raw,
        content_hash=document_hash(raw),
    )


def _ingest_pdf(path: Path, raw: bytes, ocr: OCREngine) -> IngestedDocument:
    reader = PdfReader(str(path))
    pages: list[PageText] = []
    empty = 0
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            empty += 1
        pages.append(PageText(page=index, text=text))
    ocr_used = False
    confidence = 1.0
    warnings: list[str] = []
    if pages and empty / max(len(pages), 1) >= 0.8:
        try:
            ocr_pages = ocr.extract_pages(path)
            pages = [PageText(page=i + 1, text=t) for i, t in enumerate(ocr_pages)]
            ocr_used = True
            confidence = 0.55
            warnings.append("OCR fallback used; extraction confidence is reduced.")
        except Exception as exc:  # noqa: BLE001
            confidence = 0.2
            warnings.append(f"PDF text extraction looks empty and OCR failed: {exc}")
    return IngestedDocument(
        filename=path.name,
        file_type="pdf",
        pages=pages,
        raw_bytes=raw,
        content_hash=document_hash(raw),
        ocr_used=ocr_used,
        extraction_confidence=confidence,
        warnings=warnings,
    )


def _split_page_markers(text: str) -> list[PageText]:
    """Honor explicit 'Page N:' markers used in sample/evaluation contracts."""
    import re

    matches = list(re.finditer(r"(?im)^(?:---+)?\s*page\s+(\d+)\s*:?\s*(?:---+)?$", text))
    if not matches:
        return [PageText(page=1, text=text.strip())]
    pages: list[PageText] = []
    for i, match in enumerate(matches):
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        page_no = int(match.group(1))
        pages.append(PageText(page=page_no, text=text[start:end].strip()))
    return pages or [PageText(page=1, text=text.strip())]
